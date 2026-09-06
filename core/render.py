"""L7b — RenderPlan → .docx/.pptx/.xlsx. Deterministic. Fails closed.

The LLM never writes binary office files. It emits a schema-constrained RenderPlan;
validate_plan() resolves every evidence_ref against the KB; only then does render()
write from a template the organisation owns. A field without resolvable evidence is a
validation error, not a "best effort" render — see validate_plan()'s docstring.
"""
from __future__ import annotations

import hashlib
import os
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from docxtpl import DocxTemplate
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt

from core.config import SETTINGS
from core.kb import get_span
from core.schemas import EvidenceSpan, RenderPlan, WorkbenchError

_TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"
_DOCX_TEMPLATE = _TEMPLATES_DIR / "approval_note_v1.docx"
_PPTX_TEMPLATE = _TEMPLATES_DIR / "board_deck_v1.pptx"

_SEVERITY_COLOR = {
    "info": PptxRGBColor(0x60, 0x9E, 0xE8),
    "observation": PptxRGBColor(0x3E, 0x7C, 0xB8),
    "concern": PptxRGBColor(0xE8, 0x9A, 0x3C),
    "critical": PptxRGBColor(0xC0, 0x39, 0x2B),
}
_AMBER_FILL = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
_REVIEW_MARKER = " [⚠ low-confidence OCR — verify against original]"


class EvidenceResolutionError(WorkbenchError):
    """Raised when a RenderPlan cites a span that does not exist, or a needs_review span
    is used in a numeric field without explicit human confirmation."""


# ─────────────────────────── validation (the gate) ─────────────────────────────────


def validate_plan(plan: RenderPlan, *, allow_review_spans: bool = False) -> dict[str, EvidenceSpan]:
    """THE GATE. Runs before any file is touched.

    For every GroundedValue in plan.fields:
      - if computed_from is set: every id in it must resolve via core.kb.get_span,
        and plan must carry a non-empty `calculation` string.
      - else: evidence_ref is REQUIRED and must resolve.
    For every Finding: every id in evidence_refs must resolve (min_length=1 already enforced).
    If a resolved span has needs_review=True and the value is numeric and
    allow_review_spans is False -> EvidenceResolutionError naming the field and span.

    Returns {span_id: EvidenceSpan} for the renderer to use for footnotes.
    Raises on the FIRST failure with a message naming the field, the ref, and the reason.
    """
    resolved: dict[str, EvidenceSpan] = {}

    def resolve(ref: str, context: str) -> EvidenceSpan:
        span = resolved.get(ref)
        if span is not None:
            return span
        try:
            span = get_span(ref)
        except WorkbenchError as exc:
            raise EvidenceResolutionError(
                f"{context}: evidence_ref {ref!r} does not resolve: {exc}"
            ) from exc
        resolved[ref] = span
        return span

    for field_name, gv in plan.fields.items():
        is_numeric = isinstance(gv.value, (int, float)) and not isinstance(gv.value, bool)
        field_spans: list[EvidenceSpan] = []

        if gv.computed_from:
            if not gv.calculation:
                raise EvidenceResolutionError(
                    f"field {field_name!r}: computed_from is set but calculation is empty"
                )
            for ref in gv.computed_from:
                field_spans.append(resolve(ref, f"field {field_name!r} (computed_from)"))
        else:
            if not gv.evidence_ref:
                raise EvidenceResolutionError(
                    f"field {field_name!r}: no evidence_ref and no computed_from — "
                    "every asserted value must cite evidence"
                )
            field_spans.append(resolve(gv.evidence_ref, f"field {field_name!r}"))

        for span in field_spans:
            if span.needs_review and is_numeric and not allow_review_spans:
                raise EvidenceResolutionError(
                    f"field {field_name!r}: span {span.span_id!r} has needs_review=True "
                    "(low-confidence OCR) and the field's value is numeric — pass "
                    "allow_review_spans=True to render it anyway"
                )

    for i, finding in enumerate(plan.findings):
        for ref in finding.evidence_refs:
            resolve(ref, f"finding[{i}] {finding.title!r}")

    return resolved


# ─────────────────────────── context building (shared by all 3 renderers) ─────────


def _source_note(refs: list[str], spans: dict[str, EvidenceSpan]) -> str:
    seen: dict[str, None] = {}
    for ref in refs:
        span = spans.get(ref)
        if span is not None:
            seen[f"Doc {span.doc_id}, p.{span.page}"] = None
    return "; ".join(seen)


def _review_marker(refs: list[str], spans: dict[str, EvidenceSpan]) -> str:
    for ref in refs:
        span = spans.get(ref)
        if span is not None and span.needs_review:
            return _REVIEW_MARKER
    return ""


def _field_refs(gv) -> list[str]:
    if gv.computed_from:
        return list(gv.computed_from)
    return [gv.evidence_ref] if gv.evidence_ref else []


def _build_calculation(name: str, gv, spans: dict[str, EvidenceSpan]) -> dict:
    """L4 already ran core.sandbox.verify_calculation() before building the plan — this
    renderer never evaluates anything, it only renders the expression/steps that came
    back in `calculation` and flags disagreement if the text says so."""
    text = gv.calculation or ""
    steps = [line.strip() for line in text.splitlines() if line.strip()]
    disagreed = "disagree" in text.lower()
    return {
        "label": name,
        "expression": text,
        "steps": steps,
        "result": gv.value,
        "unit": gv.unit or "",
        "disagreed": disagreed,
        "verdict_text": "DISAGREES — HUMAN REVIEW REQUIRED" if disagreed else "AGREES",
    }


def _build_context(plan: RenderPlan, spans: dict[str, EvidenceSpan], task_id: str) -> dict:
    fields_ctx: dict[str, dict] = {}
    calculations: list[dict] = []
    for name, gv in plan.fields.items():
        refs = _field_refs(gv)
        fields_ctx[name] = {
            "value": gv.value,
            "unit": gv.unit or "",
            "source_note": _source_note(refs, spans) + _review_marker(refs, spans),
        }
        if gv.computed_from:
            calculations.append(_build_calculation(name, gv, spans))

    findings_ctx = [
        {
            "title": f.title,
            "detail": f.detail,
            "severity": f.severity,
            "source_note": _source_note(f.evidence_refs, spans),
        }
        for f in plan.findings
    ]

    provenance = [spans[span_id] for span_id in sorted(spans.keys())]

    return {
        "task_id": task_id,
        "generated_on": datetime.now(timezone.utc).strftime("%Y-%m-%d"),
        "title": plan.title,
        "recommendation": plan.recommendation,
        "prepared_by": plan.prepared_by,
        "fields": fields_ctx,
        "findings": findings_ctx,
        "calculations": calculations,
        "provenance": provenance,
    }


# ─────────────────────────── DOCX ───────────────────────────────────────────────────


def _render_docx(plan: RenderPlan, spans: dict[str, EvidenceSpan], out: Path) -> None:
    if not _DOCX_TEMPLATE.is_file():
        raise WorkbenchError(f"missing docx template: {_DOCX_TEMPLATE}")

    context = _build_context(plan, spans, out.parent.name)

    tpl = DocxTemplate(str(_DOCX_TEMPLATE))
    tpl.render(context)
    tpl.save(str(out))


# ─────────────────────────── PPTX ───────────────────────────────────────────────────


def _add_textbox(slide, text: str, *, left, top, width, height, size=PptxPt(14)) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.splitlines() or [""]
    tf.text = lines[0]
    tf.paragraphs[0].font.size = size
    for line in lines[1:]:
        para = tf.add_paragraph()
        para.text = line
        para.font.size = size


def _render_pptx(plan: RenderPlan, spans: dict[str, EvidenceSpan], out: Path) -> None:
    if not _PPTX_TEMPLATE.is_file():
        raise WorkbenchError(f"missing pptx template: {_PPTX_TEMPLATE}")

    context = _build_context(plan, spans, out.parent.name)
    prs = Presentation(str(_PPTX_TEMPLATE))
    layout = prs.slide_layouts[1]  # "Title and Content" — the one built-in layout used throughout

    # 1. Title slide
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = plan.title
    _add_textbox(
        slide,
        f"Date: {context['generated_on']}\nPrepared by: {plan.prepared_by}",
        left=PptxInches(0.5), top=PptxInches(2.0), width=PptxInches(9), height=PptxInches(1.5),
    )

    # 2. Summary slide
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Summary"
    _add_textbox(
        slide, plan.recommendation,
        left=PptxInches(0.5), top=PptxInches(1.5), width=PptxInches(9), height=PptxInches(4),
    )

    # 3. One slide per Finding
    for finding in context["findings"]:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"[{finding['severity'].upper()}] {finding['title']}"
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, PptxInches(0.3), PptxInches(1.4), PptxInches(0.15), PptxInches(4)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = _SEVERITY_COLOR.get(finding["severity"], _SEVERITY_COLOR["info"])
        bar.line.fill.background()
        _add_textbox(
            slide, f"{finding['detail']}\n\nSource: {finding['source_note']}",
            left=PptxInches(0.7), top=PptxInches(1.5), width=PptxInches(8.5), height=PptxInches(4),
        )

    # 4. Values table slide
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Measured Values"
    rows = len(context["fields"]) + 1
    table_shape = slide.shapes.add_table(rows, 3, PptxInches(0.5), PptxInches(1.5), PptxInches(9), PptxInches(0.4 * rows))
    table = table_shape.table
    for col, header in enumerate(("Name", "Value", "Source")):
        table.cell(0, col).text = header
    for r, (name, f) in enumerate(context["fields"].items(), start=1):
        table.cell(r, 0).text = name
        table.cell(r, 1).text = f"{f['value']} {f['unit']}".strip()
        table.cell(r, 2).text = f["source_note"]

    # 5. Provenance slide
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Provenance"
    lines = [
        f"[{s.span_id}] {s.doc_id} p.{s.page} — \"{s.text[:80]}\""
        for s in context["provenance"]
    ]
    _add_textbox(
        slide, "\n".join(lines) or "(no cited spans)",
        left=PptxInches(0.3), top=PptxInches(1.3), width=PptxInches(9.4), height=PptxInches(5.5),
        size=PptxPt(8),
    )

    prs.save(str(out))


# ─────────────────────────── XLSX ───────────────────────────────────────────────────


def _autosize(ws) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[letter].width = min(max(length + 2, 10), 80)


def _render_xlsx(plan: RenderPlan, spans: dict[str, EvidenceSpan], out: Path) -> None:
    context = _build_context(plan, spans, out.parent.name)
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append(["Title", plan.title])
    summary.append(["Recommendation", plan.recommendation])
    summary.append(["Generated", context["generated_on"]])
    summary.append(["Prepared by", plan.prepared_by])
    for row in summary.iter_rows(min_row=1, max_row=summary.max_row, min_col=1, max_col=1):
        row[0].font = Font(bold=True)
    _autosize(summary)

    values = wb.create_sheet("Values")
    header = ["Name", "Value", "Unit", "Evidence span_id", "Doc", "Page", "Needs Review"]
    values.append(header)
    for cell in values[1]:
        cell.font = Font(bold=True)
    values.freeze_panes = "A2"
    for name, gv in plan.fields.items():
        refs = _field_refs(gv)
        first_span = spans.get(refs[0]) if refs else None
        needs_review = any(spans[r].needs_review for r in refs if r in spans)
        row_idx = values.max_row + 1
        values.append([
            name,
            gv.value,
            gv.unit or "",
            ", ".join(refs),
            first_span.doc_id if first_span else "",
            first_span.page if first_span else "",
            "YES" if needs_review else "",
        ])
        if needs_review:
            for cell in values[row_idx]:
                cell.fill = _AMBER_FILL
    _autosize(values)

    calcs = wb.create_sheet("Calculations")
    calc_header = ["Calculation", "Expression", "Step", "Result", "Unit", "Verdict"]
    calcs.append(calc_header)
    for cell in calcs[1]:
        cell.font = Font(bold=True)
    calcs.freeze_panes = "A2"
    for calc in context["calculations"]:
        steps = calc["steps"] or [""]
        for step in steps:
            row_idx = calcs.max_row + 1
            calcs.append([calc["label"], calc["expression"], step, calc["result"], calc["unit"], calc["verdict_text"]])
            if calc["disagreed"]:
                for cell in calcs[row_idx]:
                    cell.fill = _AMBER_FILL
    _autosize(calcs)

    wb.save(str(out))


# ─────────────────────────── entry point ────────────────────────────────────────────

_RENDERERS = {"docx": _render_docx, "pptx": _render_pptx, "xlsx": _render_xlsx}


def _slugify(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return slug or "deliverable"


def render(plan: RenderPlan, task_id: str, *, allow_review_spans: bool = False) -> Path:
    """validate_plan() -> dispatch on plan.output_format -> write to
    SETTINGS.paths.outputs / task_id / <slug>.<ext>. Returns the path.
    Never writes a partial file: render to a temp path, then atomic rename."""
    spans = validate_plan(plan, allow_review_spans=allow_review_spans)

    renderer = _RENDERERS.get(plan.output_format)
    if renderer is None:
        raise WorkbenchError(f"unsupported output_format: {plan.output_format!r}")

    out_dir = SETTINGS.paths.outputs / task_id
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{_slugify(plan.title)}.{plan.output_format}"

    fd, tmp_name = tempfile.mkstemp(dir=out_dir, suffix=f".{plan.output_format}.tmp")
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        renderer(plan, spans, tmp_path)
        os.replace(tmp_path, out_path)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise
    return out_path


def file_sha256(path: Path) -> str:
    """Consumed by L8 for artifact_hashes."""
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()
