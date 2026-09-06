"""L5 — the local tool surface. Exactly four families. Do not add a fifth."""
from __future__ import annotations

import base64
import hashlib
import json
import logging
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

import openpyxl

from core.config import SETTINGS
from core.sandbox import run_python, run_tests, workspace_for
from core.schemas import ToolCall, WorkbenchError

logger = logging.getLogger(__name__)


def _resolve_path(task_id: str, raw: str, *, for_write: bool) -> tuple[Path, bool]:
    """Resolve `raw` and enforce the path jail.

    Reads must land inside the task workspace or the read-only uploads directory; anything
    else — including any relative path using '..' to climb out — is a hard WorkbenchError,
    never a silent clamp.

    Writes get the same '..' hard-block, but a *clean* absolute path that simply lies outside
    the workspace is not blocked outright: it comes back with requires_approval=True so the
    caller can write nothing and surface the request to a human instead.
    """
    ws = workspace_for(task_id).resolve()
    uploads = SETTINGS.paths.uploads.resolve()
    raw_path = Path(raw)

    is_traversal = not raw_path.is_absolute() and ".." in raw_path.parts
    resolved = (raw_path if raw_path.is_absolute() else ws / raw_path).resolve()

    def _inside(root: Path) -> bool:
        try:
            resolved.relative_to(root)
            return True
        except ValueError:
            return False

    if _inside(ws):
        return resolved, False

    if is_traversal:
        raise WorkbenchError(f"path traversal is not allowed: {raw}")

    if for_write:
        return resolved, True

    if _inside(uploads):
        return resolved, False

    raise WorkbenchError(f"path escapes the task workspace and uploads directory: {raw}")


# ─────────────────────────── fs.read / fs.write ────────────────────────────────


def fs_read(task_id: str, args: dict[str, Any]) -> str:
    path = args["path"]
    binary = args.get("binary", False)
    resolved, _ = _resolve_path(task_id, path, for_write=False)
    if not resolved.is_file():
        raise WorkbenchError(f"no such file: {path}")
    data = resolved.read_bytes()
    return base64.b64encode(data).decode("ascii") if binary else data.decode("utf-8")


def fs_write(task_id: str, args: dict[str, Any]) -> dict[str, Any]:
    path = args["path"]
    content = args["content"]
    binary = args.get("binary", False)
    resolved, requires_approval = _resolve_path(task_id, path, for_write=True)
    if requires_approval:
        return {"path": str(resolved), "requires_approval": True, "written": False}
    resolved.parent.mkdir(parents=True, exist_ok=True)
    resolved.write_bytes(base64.b64decode(content)) if binary else resolved.write_text(content)
    return {"path": str(resolved), "requires_approval": False, "written": True}


# ─────────────────────────── code.exec / code.test ─────────────────────────────


def code_exec(task_id: str, args: dict[str, Any]) -> dict[str, Any]:
    filename = args.get("filename", "main.py")
    input_files_b64 = args.get("input_files") or {}
    input_files = {name: base64.b64decode(data) for name, data in input_files_b64.items()}
    result = run_python(
        task_id,
        args["code"],
        filename=filename,
        input_files=input_files or None,
        timeout_s=args.get("timeout_s"),
    )
    return result.model_dump()


def code_test(task_id: str, args: dict[str, Any]) -> dict[str, Any]:
    result = run_tests(task_id, args["code"], args["tests"])
    return result.model_dump()


# ─────────────────────────── sheet.read / sheet.write ──────────────────────────


def sheet_read(task_id: str, args: dict[str, Any]) -> dict[str, Any]:
    path = args["path"]
    sheet_name = args.get("sheet_name")
    resolved, _ = _resolve_path(task_id, path, for_write=False)
    if not resolved.is_file():
        raise WorkbenchError(f"no such spreadsheet: {path}")
    wb = openpyxl.load_workbook(resolved, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        title = ws.title
    finally:
        wb.close()
    return {"sheet_name": title, "rows": rows}


def sheet_write(task_id: str, args: dict[str, Any]) -> dict[str, Any]:
    path = args["path"]
    rows = args["rows"]
    sheet_name = args.get("sheet_name", "Sheet1")
    resolved, requires_approval = _resolve_path(task_id, path, for_write=True)
    if requires_approval:
        return {"path": str(resolved), "requires_approval": True, "written": False}
    resolved.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    wb.save(resolved)
    return {"path": str(resolved), "requires_approval": False, "written": True}


# ─────────────────────────── kb.search ──────────────────────────────────────────


def kb_search(task_id: str, args: dict[str, Any]) -> Any:
    """Delegates to core.kb — retrieval is not implemented here (L6's job)."""
    try:
        from core.kb import search as _kb_search
    except ImportError as e:
        raise WorkbenchError(f"core.kb is not available: {e}") from e
    return _kb_search(**args)


# ─────────────────────────── registry & entry point ────────────────────────────

TOOL_REGISTRY: dict[str, Callable[[str, dict[str, Any]], Any]] = {
    "fs.read": fs_read,
    "fs.write": fs_write,
    "code.exec": code_exec,
    "code.test": code_test,
    "sheet.read": sheet_read,
    "sheet.write": sheet_write,
    "kb.search": kb_search,
}


def call_tool(task_id: str, step_id: str, name: str, args: dict[str, Any]) -> tuple[Any, ToolCall]:
    """Single entry point. Validates the name, times the call, hashes args, builds the
    ToolCall record for the L8 receipt, and re-raises WorkbenchError on failure.

    On failure the ToolCall record is still built (ok=False) and attached to the raised
    WorkbenchError as `.tool_call`, so a caller that catches it can still log it to the trace.
    """
    if name not in TOOL_REGISTRY:
        raise WorkbenchError(f"unknown tool: {name}")

    call_id = f"{step_id}.{uuid.uuid4().hex[:8]}"
    args_sha256 = hashlib.sha256(json.dumps(args, sort_keys=True, default=str).encode("utf-8")).hexdigest()
    started_at = datetime.now(timezone.utc)
    t0 = time.monotonic()

    try:
        result = TOOL_REGISTRY[name](task_id, args)
        duration_ms = (time.monotonic() - t0) * 1000
        tool_call = ToolCall(
            call_id=call_id,
            step_id=step_id,
            tool_name=name,
            args_sha256=args_sha256,
            started_at=started_at,
            duration_ms=duration_ms,
            ok=True,
            error=None,
        )
        return result, tool_call
    except Exception as e:
        duration_ms = (time.monotonic() - t0) * 1000
        tool_call = ToolCall(
            call_id=call_id,
            step_id=step_id,
            tool_name=name,
            args_sha256=args_sha256,
            started_at=started_at,
            duration_ms=duration_ms,
            ok=False,
            error=str(e),
        )
        logger.error("tool %s failed: %s", name, e)
        err = e if isinstance(e, WorkbenchError) else WorkbenchError(str(e))
        err.tool_call = tool_call  # type: ignore[attr-defined]
        raise err from e
